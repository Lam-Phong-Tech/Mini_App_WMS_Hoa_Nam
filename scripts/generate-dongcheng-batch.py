from __future__ import annotations

import csv
import html
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT.parent / "Phân loại dữ liệu Power tools DongCheng_2026_18_tieu_chi.xlsx"
SHEET_NAME = "Dữ liệu sản phẩm"


def text(value: object) -> str:
    return str(value or "").strip()


def safe_file_name(value: str) -> str:
    return re.sub(r"[^a-zA-Z0-9._-]", "_", value)


def payload_for(code: str) -> str:
    return f"HN1|SKU={code}|ITEM={code}-001"


def build_spec(record: dict[str, object], position: int) -> dict[str, object]:
    return {
        "machine_code": text(record.get("Mã máy")),
        "machine_category": text(record.get("Chủng loại máy")),
        "features": text(record.get("Tính năng")),
        "package_accessories": text(record.get("Phụ kiện và gói hàng")),
        "power_source": text(record.get("Nguồn điện")),
        "power": text(record.get("Công suất")),
        "capacity": text(record.get("Khả năng")),
        "no_load_speed": text(record.get("Tốc độ không tải")),
        "impact_rate": text(record.get("Tần suất va đập")),
        "max_torque": text(record.get("Momen xoắn tối đa")),
        "weight": text(record.get("Trọng lượng máy")),
        "function_icons": text(record.get("Phân loại theo chức năng chung (mục lục biểu tượng)")),
        "usage": text(record.get("Công dụng")) or "Phân loại theo danh mục DongCheng 2026; dùng cho quản lý kho và quét mã Mini App.",
        "field": text(record.get("Lĩnh vực")) or "Xây dựng / cơ khí / bảo trì (phân loại tham khảo)",
        "audience": text(record.get("Đối tượng người dùng")) or "Thợ chuyên nghiệp; kỹ thuật viên; đội bảo trì",
        "excel_position": position,
        "pdf_page": record.get("Trang PDF") or "",
        "catalog_page": record.get("Trang catalog") or "",
        "source": "DongCheng 2026 catalog import from Excel",
    }


def read_batch(start: int, count: int) -> list[dict[str, object]]:
    workbook = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    worksheet = workbook[SHEET_NAME]
    headers = [text(value) for value in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    all_records: list[tuple[int, dict[str, object]]] = []

    for position, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=1):
        record = dict(zip(headers, values))
        code = text(record.get("Mã máy"))
        if code:
            all_records.append((position, record))

    selected = all_records[start - 1 : start - 1 + count]
    if len(selected) != count:
        raise RuntimeError(f"Expected {count} records from {start}, got {len(selected)}")

    payload: list[dict[str, object]] = []
    for position, record in selected:
        code = text(record.get("Mã máy"))
        name = text(record.get("Tên máy")) or f"{text(record.get('Chủng loại máy'))} – {code}"
        payload.append(
            {
                "sku_code": code,
                "sku_name": name,
                "unit": "cái",
                "sku_type": "PRODUCT",
                "is_serial_tracked": False,
                "warranty_applicable": True,
                "published_to_app": True,
                "power_or_voltage": text(record.get("Nguồn điện")),
                "specification": build_spec(record, position),
            }
        )

    return payload


def render_html(rows: list[dict[str, object]], svgs: list[str], start: int, end: int) -> str:
    articles = []
    for index, (row, svg) in enumerate(zip(rows, svgs), start=start):
        code = text(row["sku_code"])
        payload = payload_for(code)
        articles.append(
            f"""    <article>
      <div class="qr">{svg}</div>
      <div class="index">#{index:03d}</div>
      <div class="code">{html.escape(code)}</div>
      <div class="raw">{html.escape(payload)}</div>
    </article>"""
        )

    return f"""<!doctype html>
<html lang="vi">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width,initial-scale=1" />
  <title>QR test WMS Hoa Nam — batch {start}-{end}</title>
  <style>
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      background: #f3f6fa;
      color: #06142a;
      font-family: Arial, system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
    }}
    header {{
      position: sticky;
      top: 0;
      z-index: 1;
      background: rgba(255,255,255,.94);
      border-bottom: 1px solid #d6e0ec;
      padding: 16px;
      backdrop-filter: blur(10px);
    }}
    h1 {{ margin: 0; font-size: 20px; letter-spacing: -.04em; }}
    p {{ margin: 6px 0 0; color: #69758a; font-size: 13px; line-height: 1.5; }}
    main {{
      display: grid;
      grid-template-columns: repeat(auto-fill, minmax(180px, 1fr));
      gap: 14px;
      padding: 16px;
    }}
    article {{
      break-inside: avoid;
      background: white;
      border: 1px solid #d6e0ec;
      border-radius: 18px;
      padding: 14px;
      box-shadow: 0 10px 24px rgba(15, 23, 42, .06);
      text-align: center;
    }}
    .qr {{
      display: grid;
      place-items: center;
      width: 100%;
      aspect-ratio: 1;
      border-radius: 14px;
      background: #fff;
      border: 1px solid #eef2f7;
      padding: 8px;
    }}
    .qr svg {{ width: 100%; height: 100%; display: block; }}
    .index {{ margin-top: 10px; color: #1c70f2; font-size: 12px; font-weight: 900; }}
    .code {{ margin-top: 4px; font-size: 15px; font-weight: 900; letter-spacing: .02em; word-break: break-word; }}
    .raw {{ margin-top: 4px; color: #69758a; font-size: 10px; font-weight: 700; word-break: break-all; }}
    @media print {{
      body {{ background: white; }}
      header {{ position: static; }}
      main {{ grid-template-columns: repeat(3, 1fr); gap: 10px; padding: 10px; }}
      article {{ box-shadow: none; border-radius: 10px; padding: 10px; }}
    }}
  </style>
</head>
<body>
  <header>
    <h1>QR test WMS Hoa Nam — batch {start}-{end}</h1>
    <p>30 mã kế tiếp lấy theo cột “Mã máy” trong Excel. Payload composite: HN1|SKU=&lt;mã&gt;|ITEM=&lt;mã&gt;-001.</p>
  </header>
  <main>
{chr(10).join(articles)}
  </main>
</body>
</html>
"""


def make_qr_svg(payload: str) -> str:
    js = (
        "const QRCode=require('qrcode');"
        f"QRCode.toString({json.dumps(payload)},{{type:'svg',errorCorrectionLevel:'M',margin:2,color:{{dark:'#06142a',light:'#ffffff'}}}})"
        ".then(x=>process.stdout.write(x));"
    )
    return subprocess.check_output(["node", "-e", js], cwd=ROOT, text=True, encoding="utf-8")


def write_batch(start: int, count: int) -> None:
    end = start + count - 1
    rows = read_batch(start, count)
    output_dir = ROOT / "docs" / f"qr-test-codes-composite-{start}-{end}"
    svg_dir = output_dir / "svg"
    import_dir = ROOT / "docs" / "dongcheng-import"
    svg_dir.mkdir(parents=True, exist_ok=True)
    import_dir.mkdir(parents=True, exist_ok=True)

    svgs: list[str] = []
    codes: list[str] = []
    payloads: list[str] = []

    for index, row in enumerate(rows, start=start):
        code = text(row["sku_code"])
        payload = payload_for(code)
        svg = make_qr_svg(payload)
        (svg_dir / f"{index:03d}-{safe_file_name(code)}.svg").write_text(svg, encoding="utf-8")
        codes.append(code)
        payloads.append(payload)
        svgs.append(svg)

    (output_dir / "codes.txt").write_text("\n".join(codes) + "\n", encoding="utf-8")
    (output_dir / "payloads.txt").write_text("\n".join(payloads) + "\n", encoding="utf-8")
    (output_dir / f"qr-test-codes-composite-{start}-{end}.html").write_text(
        render_html(rows, svgs, start, end),
        encoding="utf-8",
    )

    json_path = import_dir / f"dongcheng-{start}-{end}-sku-payload.json"
    csv_path = import_dir / f"dongcheng-{start}-{end}-sku-import.csv"
    json_path.write_text(json.dumps(rows, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    with csv_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=[
                "sku_code",
                "sku_name",
                "unit",
                "sku_type",
                "is_serial_tracked",
                "warranty_applicable",
                "published_to_app",
                "power_or_voltage",
                "specification",
            ],
        )
        writer.writeheader()
        for row in rows:
            csv_row = dict(row)
            csv_row["specification"] = json.dumps(csv_row["specification"], ensure_ascii=False)
            writer.writerow(csv_row)

    print(json.dumps({
        "output_dir": str(output_dir),
        "html": str(output_dir / f"qr-test-codes-composite-{start}-{end}.html"),
        "json": str(json_path),
        "csv": str(csv_path),
        "count": len(rows),
        "first": codes[0],
        "last": codes[-1],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    batch_start = int(sys.argv[1]) if len(sys.argv) > 1 else 91
    batch_count = int(sys.argv[2]) if len(sys.argv) > 2 else 30
    write_batch(batch_start, batch_count)
